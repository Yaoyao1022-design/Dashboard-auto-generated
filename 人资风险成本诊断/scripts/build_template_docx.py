# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = "/Users/lizhiyao/Desktop/Dashboard自动化生成/人资风险成本诊断/数据看板需求填写模版.docx"

BLUE = RGBColor(0x3C, 0x6E, 0xF0)
GREY = RGBColor(0x52, 0x57, 0x65)
BLACK = RGBColor(0x23, 0x25, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "3C6EF0"
ROW_BG = "F5F5F6"
EXAMPLE_BG = "F8FAFF"


def set_run_font(run, size=11, bold=False, color=BLACK, name="PingFang SC"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element.get_or_add_rPr()
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "E4E5E9")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text, bold=False, color=BLACK, size=10, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = 1


def add_table(doc, headers, rows, header=True, example_first=False):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, fill=HEADER_BG, center=True)
    for r_idx, row in enumerate(rows):
        fill = EXAMPLE_BG if example_first and r_idx == 0 else None
        if not fill and r_idx % 2 == 1:
            fill = ROW_BG
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val or " ", size=10, fill=fill)
    doc.add_paragraph()
    return table


def add_title(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else BLACK)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)


def add_p(doc, text, bold=False, color=BLACK, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_mixed(doc, parts, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)
    return p


def blank_lines(doc, n=3, hint="（在此填写）"):
    p = doc.add_paragraph()
    run = p.add_run(hint)
    set_run_font(run, size=10, color=GREY)
    for _ in range(n):
        doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    styles = doc.styles["Normal"]
    styles.font.name = "PingFang SC"
    styles.font.size = Pt(11)
    styles.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("数据看板需求填写模版")
    set_run_font(run, size=22, bold=True, color=BLUE)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("给产品、设计、业务填写  ·  填完后交给 AI 搭看板")
    set_run_font(run, size=11, color=GREY)

    add_p(doc, "请用大白话写清楚三件事：")
    for i, t in enumerate(
        [
            "有哪些页面，用户从上到下能看到什么",
            "每个数字、表格、排行里有哪些字段",
            "点哪里会跳转、下钻或导出",
        ],
        1,
    ):
        add_p(doc, "%d. %s" % (i, t))
    add_p(doc, "能打勾就打勾，能填表就填表。没有的就写「无」或划掉，不要空着让人猜。", color=GREY)

    add_title(doc, "填写前先看这 6 条", 1)
    rules = [
        "先勾本期要做的页面，没勾的后面写「本期不做」。",
        "页面里每一块都要写：什么时候出现、怎么摆、有哪些字段、点了会怎样。",
        "字段请写中文名，并尽量补一个英文代号（方便对数据），例如：组织层级 orgLevel。",
        "有「仅部分情况才显示」的，写清楚条件，例如：选「条线」才出现条线下拉。",
        "每个表或排行至少写 1～2 行真实样子的示例（可以是假数据）。",
        "不要写开发组件名，写业务样子即可，例如「顶部两张大卡切换」。",
    ]
    for i, t in enumerate(rules, 1):
        add_p(doc, "%d. %s" % (i, t))

    add_title(doc, "0. 这是什么看板", 1)
    add_table(
        doc,
        ["项", "请填写"],
        [
            ["看板名称", ""],
            ["给谁用", ""],
            ["本期要做的页面", "☐ 驾驶舱（一个页面里切多个主题）  ☐ 独立大屏  ☐ 任务中心  ☐ 其他：________"],
            ["版本 / 日期", ""],
        ],
    )
    add_p(doc, "用户打开后，最想一眼看清什么？（最多 3 句）", bold=True)
    add_p(doc, "1. ________________________________")
    add_p(doc, "2. ________________________________")
    add_p(doc, "3. ________________________________")

    add_title(doc, "1. 有哪些页面", 1)
    add_p(doc, "画一棵结构树，让人知道页面之间的关系。抽屉是浮在页面上的明细，一般不算单独菜单。")
    add_p(doc, "示例：", bold=True)
    for line in [
        "人资成本诊断",
        "├── 驾驶舱（可切换：人力成本 / 跑冒滴漏）",
        "├── 人力成本大屏（只有人力成本）",
        "├── 跑冒滴漏大屏（只有跑冒滴漏）",
        "└── 任务中心",
        "抽屉：指标明细、综合人工成本、AI报告、组织明细、人员明细",
    ]:
        add_p(doc, line, size=10, color=GREY)
    add_p(doc, "请在下面写你们的结构：", bold=True)
    blank_lines(doc, 5)

    add_title(doc, "1.1 页面清单", 2)
    add_table(
        doc,
        ["页面名称", "地址（可后补）", "这一页讲什么", "要筛选吗", "能打开抽屉吗", "本期做吗"],
        [
            ["例：驾驶舱", "/cockpit", "一个页里切换两个主题", "要", "能", "做"],
            ["", "", "", "", "", ""],
            ["", "", "", "", "", ""],
            ["", "", "", "", "", ""],
        ],
        example_first=True,
    )

    add_title(doc, "1.2 抽屉清单", 2)
    add_table(
        doc,
        ["抽屉名称", "从哪打开", "覆盖哪些页面", "同时只能开一个吗", "本期做吗"],
        [
            ["例：AI报告", "主题卡上的「AI报告」", "驾驶舱、大屏", "是", "做"],
            ["", "", "", "是", ""],
            ["", "", "", "是", ""],
            ["", "", "", "是", ""],
        ],
        example_first=True,
    )

    add_title(doc, "2. 顶部筛选", 1)
    add_mixed(doc, [("哪些页面共用这套筛选：", True), (" ________________________________", False)])
    add_p(doc, "查一下后，刷新哪些内容：")
    add_p(doc, "☐ 整页数字都变    ☐ 只变当前主题    ☐ 其他：________")

    add_title(doc, "2.1 筛选项", 2)
    add_table(
        doc,
        ["筛选项", "怎么选", "选项有哪些", "什么时候出现", "改了会连带怎样"],
        [
            ["例：组织层级", "单选", "物流总部 / 省区 / 条线 / C1", "一直有", "换层级后，清空省区、条线、部门"],
            ["例：月", "选月份", "如 2026年08月", "一直有", "无"],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
        ],
        example_first=True,
    )

    add_p(doc, "筛选上的按钮", bold=True)
    add_table(
        doc,
        ["按钮", "按下后做什么"],
        [["重置", "回到默认：________"], ["查询", "按当前条件刷新：________"], ["其他", ""]],
    )
    add_p(doc, "打开页面时的默认筛选", bold=True)
    add_table(doc, ["筛选项", "默认值"], [["", ""], ["", ""]])

    add_title(doc, "3. 每个页面长什么样（从上到下）", 1)
    add_p(doc, "按用户扫视顺序写。一个页面有多块，就把下面「一块看板」复制多份。")
    add_p(doc, "先选这块是什么类型（只选一个）：", bold=True)
    for t in [
        "☐ 顶部大卡切换主题",
        "☐ 指标数字（可带同比、环比）",
        "☐ 趋势图",
        "☐ 地图",
        "☐ 排行榜",
        "☐ 可展开的表格",
        "☐ 同一区域里的多个下探 Tab",
        "☐ 场景小卡",
        "☐ 人员列表 / 人员卡",
        "☐ 预算或监控卡",
        "☐ AI 摘要文案",
        "☐ 任务列表",
        "☐ 其他（先写为什么上面都不合适）：________",
    ]:
        add_p(doc, t)

    add_title(doc, "一块看板（复制后填写）", 2)
    add_mixed(
        doc,
        [
            ("这块叫：", True),
            (" ____________________    ", False),
            ("属于哪个页面：", True),
            (" ____________________", False),
        ],
    )
    add_table(
        doc,
        ["问什么", "请填写"],
        [
            ["什么时候出现", "例：一直有 / 只在驾驶舱 / 选「省区」时 / 切到「跑冒滴漏」时"],
            ["怎么摆", "例：通栏 / 左边指标右边图 / 4 列 3 行"],
            ["块里有没有切换", "无 / 有，例如：YTD 和当月"],
            ["块里有哪些按钮或链接", "无 / 例如：「查看指标详情」"],
            ["没数据时怎么办", "不显示 / 显示这句话：________"],
        ],
    )
    add_p(doc, "这块里有哪些数字或列", bold=True)
    add_table(
        doc,
        ["名称", "单位", "还要显示什么", "备注"],
        [
            ["例：综合费率", "%", "同比、目标差", "大卡主数字"],
            ["例：综合人工成本", "亿元", "同比、环比", ""],
            ["", "", "", ""],
            ["", "", "", ""],
            ["", "", "", ""],
        ],
        example_first=True,
    )
    add_p(doc, "如果是表格或排行，把列写全", bold=True)
    add_table(
        doc,
        ["列名", "是否要涨跌颜色", "说明"],
        [["例：YTD费率", "否", ""], ["例：当月同比", "是", ""], ["", "", ""], ["", "", ""]],
        example_first=True,
    )
    add_p(doc, "请写 1～2 行示例（假数据即可）", bold=True)
    add_p(doc, "例：综合费率 6.42%，同比 -5.85%，环比 +6.32%", size=10, color=GREY)
    add_p(doc, "例：排行第 1 名上海，综合费率 6.42%，同比 +0.57%", size=10, color=GREY)
    blank_lines(doc, 3)
    add_p(doc, "点了会怎样", bold=True)
    add_table(
        doc,
        ["点哪里", "动作", "然后怎样"],
        [
            ["例：主题卡", "切换", "下面整块看板换成对应主题"],
            ["例：「查看指标详情」", "点击", "打开「全量指标明细」抽屉"],
            ["", "", ""],
            ["", "", ""],
        ],
        example_first=True,
    )

    add_title(doc, "填写示例（看懂后可删）", 2)
    add_p(doc, "这块叫：主题切换卡　　属于：驾驶舱", bold=True)
    add_p(doc, "类型：顶部大卡切换主题")
    add_p(doc, "什么时候出现：只有驾驶舱")
    add_p(doc, "怎么摆：并排两张大卡，同时只能选中一张")
    add_p(doc, "按钮：「AI报告」可查看、可下载")
    add_table(
        doc,
        ["卡", "主数字", "还带什么"],
        [
            ["人力成本分析", "YTD 费率、当月费率", "同比、目标差 / 环比；一段 AI 摘要"],
            ["跑冒滴漏分析", "YTD 异常条数、当月异常条数", "挽损、闭环率；一段 AI 摘要"],
        ],
    )
    add_p(doc, "点「人力成本分析」→ 下面出人力成本看板；点「跑冒滴漏分析」→ 下面出巡检看板。")

    add_title(doc, "4. 点进去看到的明细（抽屉）", 1)
    add_p(doc, "抽屉是从看板点进去后滑出来的明细，不是新菜单。")
    add_p(doc, "打开一个抽屉时：  ☐ 关掉其他抽屉（推荐）    ☐ 可以叠好几个")
    add_title(doc, "一个抽屉（复制后填写）", 2)
    add_mixed(doc, [("抽屉名称：", True), (" ____________________", False)])
    add_table(
        doc,
        ["问什么", "请填写"],
        [
            ["从哪打开", ""],
            ["大概多宽", "例：大半屏 / 全屏旁栏 / 800 宽"],
            ["是不是可以展开的树表", "是 / 否"],
            ["第一列叫什么", "例：指标名称 / 条线 / 员工姓名"],
            ["要不要一层层往下点", "无 / 有，层级是：________"],
            ["最外层要不要汇总数字", "无 / 有：________"],
            ["底部按钮", "无 / 例：下载报告、导出当前层"],
            ["导出什么", "无 / 表格哪些列 / 报告分哪几章"],
        ],
    )
    add_p(doc, "列", bold=True)
    add_table(doc, ["列名", "要涨跌颜色吗"], [["", ""], ["", ""], ["", ""]])
    add_p(doc, "示例一行：", bold=True)
    blank_lines(doc, 2)

    add_title(doc, "5. 所有点击行为（汇总）", 1)
    add_p(doc, "前面各块写过的点击，这里再收成一张总表，避免漏。")
    add_table(
        doc,
        ["用户在哪", "做什么", "然后怎样"],
        [
            ["例：顶部大卡", "切换主题", "下方看板切换，并关掉已打开的抽屉"],
            ["例：组织排行「查看全量」", "点击", "打开组织明细抽屉，可继续往下点并导出"],
            ["", "", ""],
            ["", "", ""],
            ["", "", ""],
        ],
        example_first=True,
    )

    add_title(doc, "6. 固定选项（避免前后写法不一致）", 1)
    add_p(doc, "把会反复出现的选项集中写在这里。上面各节提到时，和这里保持一致。")
    add_table(
        doc,
        ["这组选项叫什么", "有哪些值", "用在哪"],
        [
            ["例：组织层级", "物流总部 / 省区 / 条线 / C1", "筛选条"],
            ["例：主题", "人力成本分析 / 跑冒滴漏分析", "顶部大卡"],
            ["例：指标口径", "YTD / 当月", "指标矩阵"],
            ["", "", ""],
            ["", "", ""],
        ],
        example_first=True,
    )

    add_title(doc, "7. 本期做什么、不做什么", 1)
    add_table(
        doc,
        ["项", "请填写"],
        [
            ["本期不做的页面或菜单", ""],
            ["先做个入口、点进去暂不实现的", "例：任务「查看详情」只弹一句提示"],
            ["图上的线 / 柱", "☐ 先用示意数据    ☐ 已有真实月份数据"],
            ["AI 报告正文", "☐ 先占位，下载里有摘要和三章建议    ☐ 已有完整正文"],
            ["其他先不做的", ""],
        ],
    )

    add_title(doc, "8. 提交前请勾一下", 1)
    for t in [
        "☐ 第 0 节已勾本期页面，第 7 节写了不做的",
        "☐ 每个页面都按从上到下写了区块",
        "☐ 每个区块都有字段、示例、点击后去哪",
        "☐ 筛选「什么时候出现」写清楚了",
        "☐ 第 5 节能对上所有按钮、下钻、导出",
        "☐ 没有写开发组件名",
    ]:
        add_p(doc, t)

    add_title(doc, "可选：没数据 / 没权限时", 1)
    add_p(doc, "不填则默认：没数据就不画这块，不做权限区分。", color=GREY)
    add_table(
        doc,
        ["情况", "希望怎样"],
        [["筛选后没有数据", ""], ["某人看不到某个主题或抽屉", ""], ["加载失败", ""]],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
